import io
import os
import re
from typing import List

import streamlit as st
from google import genai
from google.genai import types
from pydantic import BaseModel, Field
from pypdf import PdfReader
from docx import Document


# -----------------------------
# Page configuration
# -----------------------------
st.set_page_config(
    page_title="Resume ATS Analyzer",
    page_icon="📄",
    layout="wide",
)

st.title("📄 Resume ATS Analyzer")
st.caption(
    "Upload a resume and optionally a job description to get a simulated ATS "
    "compatibility score, keyword analysis, and actionable improvements."
)

# -----------------------------
# Structured Gemini response
# -----------------------------
class KeywordItem(BaseModel):
    keyword: str = Field(description="Important keyword or phrase.")
    found: bool = Field(description="Whether the keyword is present in the resume.")
    recommendation: str = Field(
        description="Short recommendation if the keyword is missing or weak."
    )


class Improvement(BaseModel):
    priority: str = Field(description="One of: High, Medium, Low.")
    area: str = Field(description="Improvement category.")
    issue: str = Field(description="Specific issue found in the resume.")
    action: str = Field(description="Concrete action the candidate should take.")
    example: str = Field(description="A concise example of a better wording.")


class ResumeAnalysis(BaseModel):
    ats_score: int = Field(
        ge=0,
        le=100,
        description="Simulated ATS compatibility score from 0 to 100.",
    )
    summary: str = Field(description="Short overall assessment.")
    strengths: List[str] = Field(description="Three to six resume strengths.")
    weaknesses: List[str] = Field(description="Three to six important weaknesses.")
    keyword_analysis: List[KeywordItem] = Field(
        description="Important job-specific keywords and whether they are found."
    )
    improvements: List[Improvement] = Field(
        description="Prioritized, actionable improvements."
    )
    formatting_checks: List[str] = Field(
        description="ATS-friendly formatting observations."
    )
    section_checks: List[str] = Field(
        description="Observations about standard resume sections."
    )


def extract_pdf_text(file_bytes: bytes) -> str:
    reader = PdfReader(io.BytesIO(file_bytes))
    pages = []
    for page in reader.pages:
        pages.append(page.extract_text() or "")
    return "\n".join(pages).strip()


def extract_docx_text(file_bytes: bytes) -> str:
    doc = Document(io.BytesIO(file_bytes))
    parts = []

    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text:
            parts.append(text)

    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))

    return "\n".join(parts).strip()


def extract_resume_text(uploaded_file) -> str:
    file_bytes = uploaded_file.getvalue()
    extension = uploaded_file.name.lower().split(".")[-1]

    if extension == "pdf":
        return extract_pdf_text(file_bytes)
    if extension == "docx":
        return extract_docx_text(file_bytes)

    raise ValueError("Unsupported file type. Please upload a PDF or DOCX resume.")


def get_api_key(user_api_key: str = "") -> str:
    if user_api_key.strip():
        return user_api_key.strip()

    # Prefer Streamlit Cloud secrets; fall back to local environment variable.
    try:
        if "GEMINI_API_KEY" in st.secrets:
            return st.secrets["GEMINI_API_KEY"]
    except Exception:
        pass

    return os.getenv("GEMINI_API_KEY", "")


def basic_resume_metrics(text: str) -> dict:
    words = re.findall(r"\b[\w+#.-]+\b", text)
    lower = text.lower()

    standard_sections = {
        "contact": bool(re.search(r"\b(email|phone|linkedin|github)\b", lower)),
        "summary": bool(re.search(r"\b(summary|profile|objective)\b", lower)),
        "experience": bool(re.search(r"\b(experience|employment|work history)\b", lower)),
        "education": bool(re.search(r"\beducation\b", lower)),
        "skills": bool(re.search(r"\b(skills|technical skills|core competencies)\b", lower)),
    }

    quantified = len(re.findall(r"\b\d+(?:\.\d+)?%?\b", text))
    bullet_like = len(re.findall(r"(?:^|\n)\s*(?:[-•▪◦*])\s+", text))

    return {
        "word_count": len(words),
        "standard_sections": standard_sections,
        "quantified_items": quantified,
        "bullet_like_lines": bullet_like,
        "has_email": bool(re.search(r"\b[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}\b", text, re.I)),
        "has_url": bool(re.search(r"https?://|www\.", lower)),
    }


def analyze_with_gemini(
    resume_text: str, job_description: str, user_api_key: str = ""
) -> ResumeAnalysis:
    api_key = get_api_key(user_api_key)
    if not api_key:
        raise RuntimeError(
            "Add your Gemini API key in the sidebar, or configure GEMINI_API_KEY "
            "in Streamlit secrets or as an environment variable."
        )

    client = genai.Client(api_key=api_key)

    metrics = basic_resume_metrics(resume_text)

    jd_context = (
        job_description.strip()
        if job_description.strip()
        else "No job description was provided. Evaluate against general ATS-friendly resume practices."
    )

    prompt = f"""
You are an expert resume reviewer and ATS optimization specialist.

Analyze the candidate resume below. Produce a SIMULATED ATS COMPATIBILITY SCORE,
not a claim that the resume was tested by a proprietary ATS.

Scoring guidance:
- 25 points: relevant keywords and skills
- 20 points: standard sections and information completeness
- 20 points: measurable achievements and strong action-oriented wording
- 15 points: ATS-friendly formatting and readability
- 10 points: job-title/experience alignment
- 10 points: contact information and professional links

If a job description is provided, prioritize its explicit requirements and terminology.
Do not invent candidate experience, qualifications, metrics, employers, degrees, or skills.
If something is missing, clearly say it is missing rather than assuming it exists.

Pay special attention to:
- keyword matching
- standard section headings
- action verbs
- measurable achievements
- unnecessary graphics/tables/columns when they could hurt parsing
- overly long or vague summaries
- spelling/grammar issues
- skills that are claimed but not supported by experience
- missing dates, titles, locations, or other useful context
- contact information

Return concise, practical recommendations. For examples, rewrite only wording that
is supported by the source resume; use placeholders such as [X%] when a metric is
needed but not present.

Basic extracted metrics:
{metrics}

JOB DESCRIPTION:
{jd_context}

RESUME:
{resume_text[:30000]}
"""

    response = client.models.generate_content(
        model="gemini-3.5-flash",
        contents=prompt,
        config=types.GenerateContentConfig(
            response_mime_type="application/json",
            response_schema=ResumeAnalysis.model_json_schema(),
            temperature=0.2,
        ),
    )

    if not response.text:
        raise RuntimeError("Gemini returned an empty response.")

    return ResumeAnalysis.model_validate_json(response.text)


# -----------------------------
# Sidebar
# -----------------------------
with st.sidebar:
    st.header("⚙️ Analysis Settings")
    api_key_input = st.text_input(
        "Gemini API key",
        type="password",
        placeholder="Paste your Gemini API key",
        help="Your key is used for this session and is not displayed.",
    )
    st.info(
        "For the most useful ATS score, paste the job description. "
        "Without one, the score reflects general ATS-readiness."
    )
    st.markdown("**Supported formats:** PDF, DOCX")
    st.markdown("**Model:** Gemini 3.5 Flash")

# -----------------------------
# Main input
# -----------------------------
resume_file = st.file_uploader(
    "Upload your resume",
    type=["pdf", "docx"],
    help="PDF or DOCX only.",
)

job_description = st.text_area(
    "Job description (recommended)",
    height=220,
    placeholder="Paste the complete job description here...",
)

analyze_button = st.button(
    "🔍 Analyze Resume",
    type="primary",
    use_container_width=True,
)

if analyze_button:
    if resume_file is None:
        st.warning("Please upload a PDF or DOCX resume first.")
        st.stop()

    with st.spinner("Extracting resume and analyzing it with Gemini..."):
        try:
            resume_text = extract_resume_text(resume_file)

            if not resume_text.strip():
                st.error(
                    "No readable text was extracted. If this is a scanned/image-only "
                    "PDF, convert it to a text-based PDF or DOCX first."
                )
                st.stop()

            analysis = analyze_with_gemini(
                resume_text, job_description, api_key_input
            )

        except Exception as exc:
            st.error(f"Analysis failed: {exc}")
            st.stop()

    st.success("Analysis complete.")

    # Score
    col1, col2, col3 = st.columns([1, 2, 2])
    with col1:
        st.metric("ATS Score", f"{analysis.ats_score}/100")

    with col2:
        st.progress(analysis.ats_score / 100)

    with col3:
        if analysis.ats_score >= 80:
            st.success("Strong ATS readiness")
        elif analysis.ats_score >= 60:
            st.warning("Needs improvement")
        else:
            st.error("Significant improvement recommended")

    st.markdown("### 📝 Overall Assessment")
    st.write(analysis.summary)

    tab1, tab2, tab3, tab4, tab5 = st.tabs(
        ["🎯 Keywords", "💪 Strengths", "⚠️ Weaknesses", "🛠 Improvements", "📋 Checks"]
    )

    with tab1:
        if analysis.keyword_analysis:
            for item in analysis.keyword_analysis:
                status = "✅ Found" if item.found else "❌ Missing"
                st.markdown(f"**{item.keyword}** — {status}")
                if item.recommendation:
                    st.caption(item.recommendation)
        else:
            st.info("No keyword analysis was returned.")

    with tab2:
        for item in analysis.strengths:
            st.markdown(f"- {item}")

    with tab3:
        for item in analysis.weaknesses:
            st.markdown(f"- {item}")

    with tab4:
        for item in sorted(
            analysis.improvements,
            key=lambda x: {"High": 0, "Medium": 1, "Low": 2}.get(x.priority, 3),
        ):
            badge = {"High": "🔴", "Medium": "🟠", "Low": "🟢"}.get(
                item.priority, "🔵"
            )
            st.markdown(f"#### {badge} {item.priority}: {item.area}")
            st.write(f"**Issue:** {item.issue}")
            st.write(f"**Action:** {item.action}")
            st.write(f"**Example:** {item.example}")

    with tab5:
        st.markdown("**Formatting observations**")
        for item in analysis.formatting_checks:
            st.markdown(f"- {item}")

        st.markdown("**Section observations**")
        for item in analysis.section_checks:
            st.markdown(f"- {item}")

    with st.expander("📄 Extracted resume text"):
        st.text(resume_text)

st.divider()
st.caption(
    "Important: ATS scores vary between applicant-tracking systems. This app provides "
    "an AI-assisted compatibility estimate and should not be treated as an official ATS score."
)
